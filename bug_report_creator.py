import datetime
import json
import os
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

def get_input(prompt, required=True, multiline=False):
    while True:
        if multiline:
            print(f"\n{prompt} (press Enter twice to finish):")
        else:
            print(f"{prompt}: ", end="")
        
        if multiline:
            lines = []
            while True:
                line = input()
                if not line:
                    if lines:
                        break
                else:
                    lines.append(line)
            value = "\n".join(lines)
        else:
            value = input()
        
        if value.strip() or not required:
            return value.strip()
        print("This field is required. Please provide a value.")

def create_bug_report():
    print("=" * 50)
    print("       BUG REPORT CREATOR")
    print("=" * 50)
    
    print("\nPlease provide the following bug details:\n")
    
    title = get_input("Bug Title")
    description = get_input("Description")
    steps = get_input("Steps to Reproduce", multiline=True)
    expected = get_input("Expected Behavior", multiline=True)
    actual = get_input("Actual Behavior", multiline=True)
    severity = get_input("Severity (Critical/Major/Minor/Low)")
    priority = get_input("Priority (High/Medium/Low)")
    environment = get_input("Environment (OS, Browser, etc.)", required=False)
    assignee = get_input("Assignee", required=False)
    attachments = get_input("Attachments/Notes", required=False)
    
    reporter = get_input("Your Name", required=False)
    
    report = {
        "title": title,
        "description": description,
        "steps_to_reproduce": steps,
        "expected_behavior": expected,
        "actual_behavior": actual,
        "severity": severity,
        "priority": priority,
        "environment": environment,
        "assignee": assignee,
        "attachments": attachments,
        "reporter": reporter,
        "created_at": datetime.datetime.now().isoformat()
    }
    
    return report

def generate_markdown(report):
    md = f"""# Bug Report

**Title:** {report['title']}
**Created:** {report['created_at']}
**Reporter:** {report['reporter'] or 'N/A'}

---

## Details

**Severity:** {report['severity']}
**Priority:** {report['priority']}
**Environment:** {report['environment'] or 'N/A'}
**Assignee:** {report['assignee'] or 'Unassigned'}

---

## Description

{report['description']}

---

## Steps to Reproduce

{report['steps_to_reproduce']}

---

## Expected Behavior

{report['expected_behavior']}

---

## Actual Behavior

{report['actual_behavior']}

---

## Attachments/Notes

{report['attachments'] or 'None'}
"""
    return md

def generate_text(report):
    text = f"""BUG REPORT
{'=' * 50}

Title: {report['title']}
Created: {report['created_at']}
Reporter: {report['reporter'] or 'N/A'}

{'-' * 50}

Severity: {report['severity']}
Priority: {report['priority']}
Environment: {report['environment'] or 'N/A'}
Assignee: {report['assignee'] or 'Unassigned'}

{'-' * 50}

Description:
{report['description']}

Steps to Reproduce:
{report['steps_to_reproduce']}

Expected Behavior:
{report['expected_behavior']}

Actual Behavior:
{report['actual_behavior']}

Attachments/Notes:
{report['attachments'] or 'None'}
"""
    return text

def generate_excel(report):
    wb = Workbook()
    ws = wb.active
    ws.title = "Bug Report"
    
    header_fill = PatternFill(start_color="667EEA", end_color="667EEA", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    severity = report['severity'].lower() if report['severity'] else ""
    severity_colors = {
        "critical": "FF0000",
        "major": "FF8C00",
        "high": "FF8C00",
        "medium": "FFD700",
        "minor": "FFFDD0",
        "low": "90EE90"
    }
    severity_fill = None
    for key, color in severity_colors.items():
        if key in severity:
            severity_fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
            break
    
    headers = ["Field", "Value"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = thin_border
        cell.alignment = Alignment(horizontal='center')
    
    wrap_columns = ["Steps to Reproduce", "Description", "Expected Behavior", "Actual Behavior"]
    
    report_fields = [
        ("Title", report['title']),
        ("Created", report['created_at']),
        ("Reporter", report['reporter'] or 'N/A'),
        ("Severity", report['severity']),
        ("Priority", report['priority']),
        ("Environment", report['environment'] or 'N/A'),
        ("Assignee", report['assignee'] or 'Unassigned'),
        ("Description", report['description']),
        ("Steps to Reproduce", report['steps_to_reproduce']),
        ("Expected Behavior", report['expected_behavior']),
        ("Actual Behavior", report['actual_behavior']),
        ("Attachments/Notes", report['attachments'] or 'None'),
    ]
    
    for row, (field, value) in enumerate(report_fields, 2):
        cell1 = ws.cell(row=row, column=1, value=field)
        cell2 = ws.cell(row=row, column=2, value=value)
        cell1.border = thin_border
        cell2.border = thin_border
        cell1.font = Font(bold=True)
        cell1.alignment = Alignment(horizontal='left', vertical='top')
        cell2.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
        
        if field == "Severity" and severity_fill:
            cell2.fill = severity_fill
        
        if field in wrap_columns:
            ws.row_dimensions[row].height = 60
    
    for col in ['A', 'B']:
        ws.column_dimensions[col].width = 30
    
    return wb

def main():
    while True:
        report = create_bug_report()
        
        print("\n" + "=" * 50)
        print("Choose output format:")
        print("1. Markdown (.md)")
        print("2. Plain Text (.txt)")
        print("3. JSON (.json)")
        print("4. Excel (.xlsx)")
        print("5. Save all formats")
        print("6. Create new report")
        print("7. Exit")
        print("=" * 50)
        
        choice = input("Enter choice (1-7): ").strip()
        
        timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        safe_title = "".join(c for c in report['title'] if c.isalnum() or c in " -_").strip()[:30]
        
        if choice in ["1", "5"]:
            md_content = generate_markdown(report)
            filename = f"bug_report_{safe_title}_{timestamp}.md"
            with open(filename, "w") as f:
                f.write(md_content)
            print(f"\nSaved: {filename}")
        
        if choice in ["2", "5"]:
            txt_content = generate_text(report)
            filename = f"bug_report_{safe_title}_{timestamp}.txt"
            with open(filename, "w") as f:
                f.write(txt_content)
            print(f"Saved: {filename}")
        
        if choice in ["3", "5"]:
            json_content = json.dumps(report, indent=2)
            filename = f"bug_report_{safe_title}_{timestamp}.json"
            with open(filename, "w") as f:
                f.write(json_content)
            print(f"Saved: {filename}")
        
        if choice in ["4", "5"]:
            wb = generate_excel(report)
            filename = f"bug_report_{safe_title}_{timestamp}.xlsx"
            wb.save(filename)
            print(f"Saved: {filename}")
        
        if choice == "1":
            print("\n--- Markdown Preview ---")
            print(generate_markdown(report))
        
        if choice == "2":
            print("\n--- Text Preview ---")
            print(generate_text(report))
        
        if choice == "3":
            print("\n--- JSON Preview ---")
            print(json.dumps(report, indent=2))
        
        if choice == "4":
            print("\n--- Excel file saved ---")
            print(f"bug_report_{safe_title}_{timestamp}.xlsx")
        
        if choice == "6":
            print("\n" + "=" * 50)
            print("Creating new report...")
            print("=" * 50)
            continue
        
        if choice == "7":
            print("Goodbye!")
            break
        
        if choice not in ["1", "2", "3", "4", "5"]:
            print("Invalid choice.")
            break
        
        break
    
    return 0

if __name__ == "__main__":
    sys.exit(main())