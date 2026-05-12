from flask import Flask, render_template, render_template_string, request, jsonify, send_file
import os
import json
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import re

app = Flask(__name__)

OPENAI_API_KEY = os.environ.get('OPENAI_API_KEY', '')

HTML_TEMPLATE = '''
<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>AI Bug Report Generator</title>
    <link rel="preconnect" href="https://fonts.googleapis.com">
    <link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
    <link href="https://fonts.googleapis.com/css2?family=Inter:wght@300;400;500;600;700&display=swap" rel="stylesheet">
    <style>
        :root {
            --primary: #6366f1;
            --primary-light: #818cf8;
            --primary-dark: #4f46e5;
            --accent: #06b6d4;
            --accent-light: #22d3ee;
            --success: #10b981;
            --warning: #f59e0b;
            --danger: #ef4444;
            --bg-dark: #0f172a;
            --bg-card: rgba(30, 41, 59, 0.7);
            --bg-input: rgba(15, 23, 42, 0.8);
            --text-primary: #f1f5f9;
            --text-secondary: #94a3b8;
            --border: rgba(148, 163, 184, 0.2);
            --glass-shadow: 0 8px 32px rgba(0, 0, 0, 0.4);
        }

        * { box-sizing: border-box; margin: 0; padding: 0; }

        body {
            font-family: 'Inter', -apple-system, BlinkMacSystemFont, sans-serif;
            background: var(--bg-dark);
            background-image:
                radial-gradient(ellipse at 20% 50%, rgba(99, 102, 241, 0.15) 0%, transparent 50%),
                radial-gradient(ellipse at 80% 20%, rgba(6, 182, 212, 0.15) 0%, transparent 50%),
                radial-gradient(ellipse at 50% 80%, rgba(16, 185, 129, 0.1) 0%, transparent 50%);
            min-height: 100vh;
            padding: 40px 20px;
            color: var(--text-primary);
            line-height: 1.6;
        }

        .container {
            max-width: 900px;
            margin: 0 auto;
            animation: fadeIn 0.6s ease-out;
        }

        @keyframes fadeIn {
            from { opacity: 0; transform: translateY(20px); }
            to { opacity: 1; transform: translateY(0); }
        }

        .header {
            text-align: center;
            margin-bottom: 40px;
        }

        h1 {
            font-size: 2.5rem;
            font-weight: 700;
            background: linear-gradient(135deg, var(--primary-light), var(--accent-light));
            -webkit-background-clip: text;
            -webkit-text-fill-color: transparent;
            background-clip: text;
            margin-bottom: 10px;
            letter-spacing: -0.5px;
        }

        .subtitle {
            color: var(--text-secondary);
            font-size: 1rem;
            font-weight: 300;
        }

        .card {
            background: var(--bg-card);
            backdrop-filter: blur(20px);
            -webkit-backdrop-filter: blur(20px);
            border-radius: 20px;
            padding: 40px;
            box-shadow: var(--glass-shadow);
            border: 1px solid var(--border);
            transition: transform 0.3s ease, box-shadow 0.3s ease;
        }

        .card:hover {
            transform: translateY(-2px);
            box-shadow: 0 12px 40px rgba(0, 0, 0, 0.5);
        }

        .form-group { margin-bottom: 28px; }

        label {
            display: block;
            margin-bottom: 12px;
            font-weight: 500;
            color: var(--text-secondary);
            font-size: 0.9rem;
            text-transform: uppercase;
            letter-spacing: 1px;
        }

        textarea {
            width: 100%;
            padding: 18px 20px;
            border: 2px solid var(--border);
            border-radius: 12px;
            font-size: 1rem;
            background: var(--bg-input);
            color: var(--text-primary);
            transition: all 0.3s ease;
            resize: vertical;
            min-height: 180px;
            font-family: 'Inter', sans-serif;
            line-height: 1.6;
        }

        textarea:focus {
            outline: none;
            border-color: var(--primary);
            box-shadow: 0 0 0 4px rgba(99, 102, 241, 0.2);
            background: rgba(15, 23, 42, 0.95);
        }

        textarea::placeholder { color: var(--text-secondary); opacity: 0.6; }

        .btn {
            background: linear-gradient(135deg, var(--primary), var(--primary-dark));
            color: white;
            border: none;
            padding: 18px 50px;
            border-radius: 12px;
            font-size: 1.05rem;
            font-weight: 600;
            cursor: pointer;
            transition: all 0.3s ease;
            width: 100%;
            margin-top: 10px;
            position: relative;
            overflow: hidden;
            letter-spacing: 0.5px;
        }

        .btn::before {
            content: '';
            position: absolute;
            top: 0;
            left: -100%;
            width: 100%;
            height: 100%;
            background: linear-gradient(90deg, transparent, rgba(255, 255, 255, 0.2), transparent);
            transition: left 0.5s ease;
        }

        .btn:hover::before { left: 100%; }

        .btn:hover {
            transform: translateY(-2px);
            box-shadow: 0 10px 30px rgba(99, 102, 241, 0.4);
        }

        .btn:active { transform: translateY(0); }

        .btn:disabled {
            background: linear-gradient(135deg, #475569, #334155);
            cursor: not-allowed;
            transform: none;
            box-shadow: none;
        }

        .btn:disabled::before { display: none; }

        .warning-box {
            background: rgba(239, 68, 68, 0.1);
            border: 1px solid rgba(239, 68, 68, 0.3);
            border-radius: 12px;
            padding: 24px;
            margin-top: 28px;
            display: none;
            animation: slideIn 0.3s ease-out;
        }

        .warning-box.show { display: block; }

        @keyframes slideIn {
            from { opacity: 0; transform: translateY(-10px); }
            to { opacity: 1; transform: translateY(0); }
        }

        .warning-box h3 {
            color: var(--danger);
            margin-bottom: 12px;
            font-size: 1.1rem;
            display: flex;
            align-items: center;
            gap: 8px;
        }

        .warning-box p { color: #fca5a5; line-height: 1.7; }

        .result-box {
            background: var(--bg-input);
            border-radius: 16px;
            padding: 32px;
            margin-top: 28px;
            display: none;
            border: 1px solid rgba(6, 182, 212, 0.3);
            animation: slideIn 0.4s ease-out;
        }

        .result-box.show { display: block; }

        .result-box h3 {
            color: var(--accent-light);
            margin-bottom: 24px;
            font-size: 1.3rem;
            font-weight: 600;
            display: flex;
            align-items: center;
            gap: 10px;
        }

        .result-box h3::before {
            content: '✓';
            display: inline-flex;
            align-items: center;
            justify-content: center;
            width: 28px;
            height: 28px;
            background: rgba(6, 182, 212, 0.2);
            border-radius: 50%;
            font-size: 0.9rem;
        }

        .field {
            margin-bottom: 22px;
            padding-bottom: 22px;
            border-bottom: 1px solid var(--border);
        }

        .field:last-child { border-bottom: none; }

        .field-label {
            color: var(--text-secondary);
            font-size: 0.8rem;
            margin-bottom: 8px;
            text-transform: uppercase;
            letter-spacing: 1px;
            font-weight: 500;
            display: flex;
            align-items: center;
            gap: 10px;
        }

        .field-value {
            color: var(--text-primary);
            font-size: 1rem;
            line-height: 1.7;
            word-wrap: break-word;
        }

        .confidence {
            display: inline-flex;
            align-items: center;
            padding: 4px 14px;
            border-radius: 20px;
            font-size: 0.8rem;
            font-weight: 600;
            letter-spacing: 0.5px;
        }

        .confidence.high {
            background: rgba(16, 185, 129, 0.2);
            color: var(--success);
            border: 1px solid rgba(16, 185, 129, 0.3);
        }

        .confidence.low {
            background: rgba(239, 68, 68, 0.2);
            color: var(--danger);
            border: 1px solid rgba(239, 68, 68, 0.3);
        }

        .btn-download {
            background: linear-gradient(135deg, var(--success), #059669);
            color: white;
            border: none;
            padding: 14px 36px;
            border-radius: 10px;
            font-size: 1rem;
            font-weight: 600;
            cursor: pointer;
            margin-top: 24px;
            transition: all 0.3s ease;
            display: inline-flex;
            align-items: center;
            gap: 8px;
        }

        .btn-download:hover {
            transform: translateY(-2px);
            box-shadow: 0 8px 25px rgba(16, 185, 129, 0.4);
        }

        .btn-download:disabled {
            background: linear-gradient(135deg, #475569, #334155);
            cursor: not-allowed;
            transform: none;
            box-shadow: none;
        }

        .spinner {
            display: none;
            width: 20px;
            height: 20px;
            border: 3px solid rgba(255, 255, 255, 0.3);
            border-top-color: white;
            border-radius: 50%;
            animation: spin 0.8s linear infinite;
        }

        @keyframes spin { to { transform: rotate(360deg); } }

        .loading {
            display: flex;
            align-items: center;
            justify-content: center;
            gap: 10px;
        }

        .json-output {
            background: rgba(15, 23, 42, 0.9);
            border-radius: 12px;
            padding: 24px;
            margin-top: 28px;
            border: 1px solid var(--border);
        }

        .json-output label {
            text-transform: none;
            letter-spacing: 0;
            font-size: 0.9rem;
            margin-bottom: 16px;
            color: var(--text-secondary);
        }

        .json-output pre {
            color: var(--accent-light);
            font-family: 'Consolas', 'Monaco', monospace;
            font-size: 0.85rem;
            overflow-x: auto;
            white-space: pre-wrap;
            word-wrap: break-word;
            line-height: 1.6;
        }

        .btn-text { display: inline-block; }

        @media (max-width: 768px) {
            body { padding: 20px 16px; }
            h1 { font-size: 2rem; }
            .card { padding: 24px; }
            .btn { padding: 16px 30px; }
        }
    </style>
</head>
<body>
    <div class="container">
        <div class="header">
            <h1>AI Bug Report Generator</h1>
            <p class="subtitle">Transform free-text bug descriptions into structured reports instantly</p>
        </div>

        <div class="card">
            <form id="bugForm">
                <div class="form-group">
                    <label for="description">Bug Description *</label>
                    <textarea
                        name="description"
                        id="description"
                        required
                        placeholder="Describe the bug in your own words...

Example: When I try to login with invalid credentials, the error message is not displayed properly. The page just reloads without showing any feedback to the user."
                    ></textarea>
                </div>

                <button type="submit" class="btn" id="generateBtn">
                    <span class="btn-text" id="btnText">Generate Bug Report</span>
                    <div class="spinner" id="spinner"></div>
                </button>
            </form>

            <div class="warning-box" id="warningBox">
                <h3>⚠️ Need More Information</h3>
                <p id="warningMessage"></p>
            </div>

            <div class="result-box" id="resultBox">
                <h3>Generated Bug Report</h3>
                <div id="reportContent"></div>
                <button class="btn-download" id="downloadBtn" onclick="downloadExcel()">
                    <svg width="18" height="18" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2" stroke-linecap="round" stroke-linejoin="round"><path d="M21 15v4a2 2 0 0 1-2 2H5a2 2 0 0 1-2-2v-4"/><polyline points="7 10 12 15 17 10"/><line x1="12" y1="15" x2="12" y2="3"/></svg>
                    Download Excel
                </button>
                <div class="json-output">
                    <label>JSON Output</label>
                    <pre id="jsonOutput"></pre>
                </div>
            </div>
        </div>
    </div>

    <script>
        let currentReport = null;

        const form = document.getElementById('bugForm');
        const generateBtn = document.getElementById('generateBtn');
        const btnText = document.getElementById('btnText');
        const spinner = document.getElementById('spinner');
        const warningBox = document.getElementById('warningBox');
        const warningMessage = document.getElementById('warningMessage');
        const resultBox = document.getElementById('resultBox');
        const reportContent = document.getElementById('reportContent');
        const jsonOutput = document.getElementById('jsonOutput');
        const downloadBtn = document.getElementById('downloadBtn');

        form.addEventListener('submit', async (e) => {
            e.preventDefault();

            const description = document.getElementById('description').value.trim();
            if (!description) {
                alert('Please enter a bug description');
                return;
            }

            generateBtn.disabled = true;
            btnText.textContent = 'Processing...';
            spinner.style.display = 'inline-block';
            warningBox.classList.remove('show');
            resultBox.classList.remove('show');

            try {
                const response = await fetch('/generate', {
                    method: 'POST',
                    headers: {'Content-Type': 'application/json'},
                    body: JSON.stringify({ description })
                });

                const data = await response.json();

                spinner.style.display = 'none';
                btnText.textContent = 'Generate Bug Report';
                generateBtn.disabled = false;

                if (data.needs_clarification) {
                    warningBox.classList.add('show');
                    warningMessage.textContent = data.message;
                    return;
                }

                currentReport = data.report;
                displayReport(data.report);

            } catch (error) {
                spinner.style.display = 'none';
                btnText.textContent = 'Generate Bug Report';
                generateBtn.disabled = false;
                alert('Error: ' + error.message);
            }
        });

        function displayReport(report) {
            const confidenceClass = report.confidence >= 80 ? 'high' : 'low';

            reportContent.innerHTML = `
                <div class="field">
                    <div class="field-label">Bug Title</div>
                    <div class="field-value">${report.title}</div>
                </div>
                <div class="field">
                    <div class="field-label">Description</div>
                    <div class="field-value">${report.description}</div>
                </div>
                <div class="field">
                    <div class="field-label">Steps to Reproduce</div>
                    <div class="field-value">${report.steps_to_reproduce}</div>
                </div>
                <div class="field">
                    <div class="field-label">Expected Behavior</div>
                    <div class="field-value">${report.expected_behavior}</div>
                </div>
                <div class="field">
                    <div class="field-label">Actual Behavior</div>
                    <div class="field-value">${report.actual_behavior}</div>
                </div>
                <div class="field">
                    <div class="field-label">Severity <span class="confidence ${confidenceClass}">${report.severity}</span></div>
                    <div class="field-value">${report.severity}</div>
                </div>
                <div class="field">
                    <div class="field-label">Priority</div>
                    <div class="field-value">${report.priority}</div>
                </div>
                <div class="field">
                    <div class="field-label">Environment</div>
                    <div class="field-value">${report.environment}</div>
                </div>
                <div class="field">
                    <div class="field-label">Test Type</div>
                    <div class="field-value">${report.test_type}</div>
                </div>
                <div class="field">
                    <div class="field-label">Status</div>
                    <div class="field-value">${report.status}</div>
                </div>
                <div class="field">
                    <div class="field-label">Confidence Score</div>
                    <div class="field-value">${report.confidence}%</div>
                </div>
            `;

            jsonOutput.textContent = JSON.stringify(report, null, 2);
            resultBox.classList.add('show');
        }

        async function downloadExcel() {
            if (!currentReport) return;

            try {
                const response = await fetch('/download-excel', {
                    method: 'POST',
                    headers: {'Content-Type': 'application/json'},
                    body: JSON.stringify(currentReport)
                });

                const blob = await response.blob();
                const url = window.URL.createObjectURL(blob);
                const a = document.createElement('a');
                a.href = url;
                a.download = `bug_report_${Date.now()}.xlsx`;
                document.body.appendChild(a);
                a.click();
                window.URL.revokeObjectURL(url);
                a.remove();
            } catch (error) {
                alert('Error downloading Excel: ' + error.message);
            }
        }
    </script>
</body>
</html>
'''

def validate_input(description):
    issues = []
    
    if not description or len(description.strip()) < 10:
        issues.append("Description is too short or empty")
    
    keywords = {
        'action': ['click', 'enter', 'submit', 'type', 'select', 'navigate', 'press', 'tap', 'open', 'close'],
        'result': ['error', 'wrong', 'incorrect', 'fail', 'crash', 'freeze', 'not', 'unable', 'cannot'],
        'location': ['page', 'screen', 'button', 'form', 'field', 'input', 'login', 'menu', 'modal', 'dropdown']
    }
    
    desc_lower = description.lower()
    
    has_action = any(kw in desc_lower for kw in keywords['action'])
    has_result = any(kw in desc_lower for kw in keywords['result'])
    has_location = any(kw in desc_lower for kw in keywords['location'])
    
    score = sum([has_action, has_result, has_location])
    
    if score < 2:
        if not has_action:
            issues.append("Missing information about what action was performed")
        if not has_result:
            issues.append("Missing information about what went wrong or the observed behavior")
        if not has_location:
            issues.append("Missing information about where the bug occurred (page/screen/element)")
    
    word_count = len(description.split())
    if word_count < 5:
        issues.append("Description is too brief (less than 5 words)")
    
    return issues, score

def generate_bug_report(description):
    issues, clarity_score = validate_input(description)
    
    if issues:
        return {
            "needs_clarification": True,
            "message": "I am not fully understanding the issue. Could you please provide more details such as steps, expected behavior, or actual behavior?",
            "issues": issues,
            "clarity_score": clarity_score
        }
    
    if OPENAI_API_KEY:
        return generate_with_openai(description)
    else:
        return generate_local(description)

def generate_with_openai(description):
    import requests
    
    prompt = f'''Analyze this bug description and generate a structured bug report. 

IMPORTANT RULES:
1. Only generate if description is clear with ≥80% confidence
2. If unclear, return "NEEDS_CLARIFICATION" with specific missing info
3. Do NOT fabricate information

Bug Description: {description}

Return a JSON object with these exact fields:
- title: Short clear summary (max 60 chars)
- description: Refined version of input
- steps_to_reproduce: Numbered list
- expected_behavior: What should happen
- actual_behavior: What actually happened
- severity: Low | Medium | High | Critical
- priority: Low | Medium | High
- environment: Web | Mobile | API | Desktop
- test_type: Functional | UI | Performance | Security
- status: New
- confidence: 0-100 score of how clear the description is
- missing_info: Array of any unclear aspects (empty if clear)

Respond ONLY with valid JSON.'''
    
    try:
        response = requests.post(
            'https://api.openai.com/v1/chat/completions',
            headers={
                'Authorization': f'Bearer {OPENAI_API_KEY}',
                'Content-Type': 'application/json'
            },
            json={
                'model': 'gpt-3.5-turbo',
                'messages': [{'role': 'user', 'content': prompt}],
                'temperature': 0.3
            },
            timeout=30
        )
        
        if response.status_code == 200:
            content = response.json()['choices'][0]['message']['content']
            result = json.loads(content)
            result['needs_clarification'] = result.get('confidence', 100) < 80
            if result['needs_clarification']:
                result['message'] = "I am not fully understanding the issue. Could you please provide more details such as steps, expected behavior, or actual behavior?"
            return result
        else:
            return generate_local(description)
    except Exception as e:
        print(f"OpenAI API error: {e}")
        return generate_local(description)

def generate_local(description):
    desc_lower = description.lower()
    
    severity_indicators = {
        'Critical': ['crash', 'freeze', 'data loss', 'security', 'breach', 'unauthorized', 'system down'],
        'Major': ['error', 'fail', 'not working', 'broken', 'wrong', 'incorrect'],
        'Medium': ['issue', 'problem', 'delay', 'slow', 'unexpected'],
        'Low': ['typo', 'cosmetic', 'minor', 'suggestion', 'improvement']
    }
    
    severity = 'Medium'
    for sev, indicators in severity_indicators.items():
        if any(ind in desc_lower for ind in indicators):
            severity = sev
            break
    
    priority_map = {'Critical': 'High', 'Major': 'High', 'Medium': 'Medium', 'Low': 'Low'}
    priority = priority_map[severity]
    
    environment = 'Web'
    if any(x in desc_lower for x in ['mobile', 'ios', 'android', 'app']):
        environment = 'Mobile'
    elif any(x in desc_lower for x in ['api', 'endpoint', 'rest', 'json', 'backend']):
        environment = 'API'
    elif any(x in desc_lower for x in ['desktop', 'exe', 'install']):
        environment = 'Desktop'
    
    test_type = 'Functional'
    if any(x in desc_lower for x in ['ui', 'visual', 'display', 'style', 'color', 'layout', 'button']):
        test_type = 'UI'
    elif any(x in desc_lower for x in ['performance', 'slow', 'load', 'speed', 'memory']):
        test_type = 'Performance'
    elif any(x in desc_lower for x in ['security', 'auth', 'login', 'password', 'injection']):
        test_type = 'Security'
    
    clarity_score = 80
    
    if any(x in desc_lower for x in ['step', 'then', 'after', 'before', 'first', 'next']):
        clarity_score += 10
    if any(x in desc_lower for x in ['expected', 'should', 'but', 'instead']):
        clarity_score += 5
    
    clarity_score = min(clarity_score, 95)
    
    action_words = ['click', 'enter', 'type', 'submit', 'select', 'navigate', 'press', 'open']
    found_actions = [w for w in action_words if w in desc_lower]
    
    if found_actions:
        steps = f"1. {found_actions[0].title()} on the relevant element\n2. Observe the behavior"
    else:
        steps = "1. Perform the described action\n2. Observe the behavior"
    
    words = description.split()
    title = ' '.join(words[:6])[:60]
    if len(words) > 6:
        title += '...'
    
    expected = "The system should behave correctly according to requirements"
    if 'error' in desc_lower or 'wrong' in desc_lower:
        expected = "Proper error message should be displayed"
    if 'login' in desc_lower:
        expected = "User should be able to login successfully or see appropriate error"
    
    actual = description
    
    if clarity_score < 80:
        return {
            "needs_clarification": True,
            "message": "I am not fully understanding the issue. Could you please provide more details such as steps, expected behavior, or actual behavior?",
            "confidence": clarity_score
        }
    
    return {
        "needs_clarification": False,
        "title": title,
        "description": description,
        "steps_to_reproduce": steps,
        "expected_behavior": expected,
        "actual_behavior": actual,
        "severity": severity,
        "priority": priority,
        "environment": environment,
        "test_type": test_type,
        "status": "New",
        "confidence": clarity_score
    }

def create_excel(report):
    wb = Workbook()
    ws = wb.active
    ws.title = "Bug Report"
    
    header_fill = PatternFill(start_color="00D9FF", end_color="00D9FF", fill_type="solid")
    header_font = Font(bold=True, color="000000")
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    headers = ["Title", "Description", "Steps to Reproduce", "Expected Behavior", "Actual Behavior", "Severity", "Priority", "Environment", "Test Type", "Status"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = thin_border
        cell.alignment = Alignment(horizontal='center', wrap_text=True)
    
    values = [
        report.get('title', ''),
        report.get('description', ''),
        report.get('steps_to_reproduce', ''),
        report.get('expected_behavior', ''),
        report.get('actual_behavior', ''),
        report.get('severity', ''),
        report.get('priority', ''),
        report.get('environment', ''),
        report.get('test_type', ''),
        report.get('status', 'New')
    ]
    
    for col, value in enumerate(values, 1):
        cell = ws.cell(row=2, column=col, value=value)
        cell.border = thin_border
        cell.alignment = Alignment(wrap_text=True, vertical='top')
    
    ws.column_dimensions['A'].width = 25
    ws.column_dimensions['B'].width = 35
    ws.column_dimensions['C'].width = 30
    ws.column_dimensions['D'].width = 30
    ws.column_dimensions['E'].width = 30
    ws.column_dimensions['F'].width = 12
    ws.column_dimensions['G'].width = 12
    ws.column_dimensions['H'].width = 12
    ws.column_dimensions['I'].width = 15
    ws.column_dimensions['J'].width = 10
    
    ws.row_dimensions[2].height = 100
    
    return wb

@app.route('/')
def index():
    return render_template_string(HTML_TEMPLATE)

@app.route('/generate', methods=['POST'])
def generate():
    data = request.json
    description = data.get('description', '').strip()
    
    if not description:
        return jsonify({
            "needs_clarification": True,
            "message": "Please provide a bug description"
        })
    
    result = generate_bug_report(description)
    return jsonify(result)

@app.route('/download-excel', methods=['POST'])
def download_excel():
    import io
    
    report = request.json
    
    if not report or report.get('needs_clarification'):
        return jsonify({"error": "No valid report to export"}), 400
    
    wb = create_excel(report)
    
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    
    filename = f"bug_report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    
    return send_file(
        buffer,
        as_attachment=True,
        download_name=filename,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )

if __name__ == '__main__':
    print("=" * 60)
    print("AI Bug Report Generator")
    print("=" * 60)
    print("\nTo use OpenAI API, set OPENAI_API_KEY environment variable")
    print("Otherwise, local generation will be used")
    print("\nStarting server at http://localhost:5000")
    print("=" * 60)
    app.run(debug=False, port=5000)
