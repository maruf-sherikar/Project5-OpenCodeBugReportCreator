from http.server import HTTPServer, BaseHTTPRequestHandler
import json
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import re

HTML = '''<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>AI Bug Report Generator</title>
    <link rel="preconnect" href="https://fonts.googleapis.com">
    <link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
    <link href="https://fonts.googleapis.com/css2?family=Inter:wght@300;400;500;600;700&display=swap" rel="stylesheet">
    <style>
        :root {
            --primary: #6366f1;
            --primary-light: #818cf8;
            --primary-dark: #4f46e5;
            --accent: #06b6d4;
            --accent-light: #22d3ee;
            --success: #10b981;
            --warning: #f59e0b;
            --danger: #ef4444;
            --bg-dark: #0f172a;
            --bg-card: rgba(30, 41, 59, 0.7);
            --bg-input: rgba(15, 23, 42, 0.8);
            --text-primary: #f1f5f9;
            --text-secondary: #94a3b8;
            --border: rgba(148, 163, 184, 0.2);
            --glass-shadow: 0 8px 32px rgba(0, 0, 0, 0.4);
        }

        * { box-sizing: border-box; margin: 0; padding: 0; }

        body {
            font-family: 'Inter', -apple-system, BlinkMacSystemFont, sans-serif;
            background: var(--bg-dark);
            background-image:
                radial-gradient(ellipse at 20% 50%, rgba(99, 102, 241, 0.15) 0%, transparent 50%),
                radial-gradient(ellipse at 80% 20%, rgba(6, 182, 212, 0.15) 0%, transparent 50%),
                radial-gradient(ellipse at 50% 80%, rgba(16, 185, 129, 0.1) 0%, transparent 50%);
            min-height: 100vh;
            padding: 40px 20px;
            color: var(--text-primary);
            line-height: 1.6;
        }

        .container {
            max-width: 900px;
            margin: 0 auto;
            animation: fadeIn 0.6s ease-out;
        }

        @keyframes fadeIn {
            from { opacity: 0; transform: translateY(20px); }
            to { opacity: 1; transform: translateY(0); }
        }

        .header {
            text-align: center;
            margin-bottom: 40px;
        }

        h1 {
            font-size: 2.5rem;
            font-weight: 700;
            background: linear-gradient(135deg, var(--primary-light), var(--accent-light));
            -webkit-background-clip: text;
            -webkit-text-fill-color: transparent;
            background-clip: text;
            margin-bottom: 10px;
            letter-spacing: -0.5px;
        }

        .subtitle {
            color: var(--text-secondary);
            font-size: 1rem;
            font-weight: 300;
        }

        .card {
            background: var(--bg-card);
            backdrop-filter: blur(20px);
            -webkit-backdrop-filter: blur(20px);
            border-radius: 20px;
            padding: 40px;
            box-shadow: var(--glass-shadow);
            border: 1px solid var(--border);
            transition: transform 0.3s ease, box-shadow 0.3s ease;
        }

        .card:hover {
            transform: translateY(-2px);
            box-shadow: 0 12px 40px rgba(0, 0, 0, 0.5);
        }

        .form-group { margin-bottom: 28px; }

        label {
            display: block;
            margin-bottom: 12px;
            font-weight: 500;
            color: var(--text-secondary);
            font-size: 0.9rem;
            text-transform: uppercase;
            letter-spacing: 1px;
        }

        textarea {
            width: 100%;
            padding: 18px 20px;
            border: 2px solid var(--border);
            border-radius: 12px;
            font-size: 1rem;
            background: var(--bg-input);
            color: var(--text-primary);
            transition: all 0.3s ease;
            resize: vertical;
            min-height: 180px;
            font-family: 'Inter', sans-serif;
            line-height: 1.6;
        }

        textarea:focus {
            outline: none;
            border-color: var(--primary);
            box-shadow: 0 0 0 4px rgba(99, 102, 241, 0.2);
            background: rgba(15, 23, 42, 0.95);
        }

        textarea::placeholder { color: var(--text-secondary); opacity: 0.6; }

        .btn {
            background: linear-gradient(135deg, var(--primary), var(--primary-dark));
            color: white;
            border: none;
            padding: 18px 50px;
            border-radius: 12px;
            font-size: 1.05rem;
            font-weight: 600;
            cursor: pointer;
            transition: all 0.3s ease;
            width: 100%;
            margin-top: 10px;
            position: relative;
            overflow: hidden;
            letter-spacing: 0.5px;
        }

        .btn::before {
            content: '';
            position: absolute;
            top: 0;
            left: -100%;
            width: 100%;
            height: 100%;
            background: linear-gradient(90deg, transparent, rgba(255, 255, 255, 0.2), transparent);
            transition: left 0.5s ease;
        }

        .btn:hover::before { left: 100%; }

        .btn:hover {
            transform: translateY(-2px);
            box-shadow: 0 10px 30px rgba(99, 102, 241, 0.4);
        }

        .btn:active { transform: translateY(0); }

        .warning-box {
            background: rgba(239, 68, 68, 0.1);
            border: 1px solid rgba(239, 68, 68, 0.3);
            border-radius: 12px;
            padding: 24px;
            margin-top: 28px;
            display: none;
            animation: slideIn 0.3s ease-out;
        }

        .warning-box.show { display: block; }

        @keyframes slideIn {
            from { opacity: 0; transform: translateY(-10px); }
            to { opacity: 1; transform: translateY(0); }
        }

        .warning-box h3 {
            color: var(--danger);
            margin-bottom: 12px;
            font-size: 1.1rem;
            display: flex;
            align-items: center;
            gap: 8px;
        }

        .warning-box p { color: #fca5a5; line-height: 1.7; }

        .result-box {
            background: var(--bg-input);
            border-radius: 16px;
            padding: 32px;
            margin-top: 28px;
            display: none;
            border: 1px solid rgba(6, 182, 212, 0.3);
            animation: slideIn 0.4s ease-out;
        }

        .result-box.show { display: block; }

        .result-box h3 {
            color: var(--accent-light);
            margin-bottom: 24px;
            font-size: 1.3rem;
            font-weight: 600;
            display: flex;
            align-items: center;
            gap: 10px;
        }

        .result-box h3::before {
            content: '✓';
            display: inline-flex;
            align-items: center;
            justify-content: center;
            width: 28px;
            height: 28px;
            background: rgba(6, 182, 212, 0.2);
            border-radius: 50%;
            font-size: 0.9rem;
        }

        .field {
            margin-bottom: 22px;
            padding-bottom: 22px;
            border-bottom: 1px solid var(--border);
        }

        .field:last-child { border-bottom: none; }

        .field-label {
            color: var(--text-secondary);
            font-size: 0.8rem;
            margin-bottom: 8px;
            text-transform: uppercase;
            letter-spacing: 1px;
            font-weight: 500;
        }

        .field-value {
            color: var(--text-primary);
            font-size: 1rem;
            line-height: 1.7;
            word-wrap: break-word;
        }

        .btn-download {
            background: linear-gradient(135deg, var(--success), #059669);
            color: white;
            border: none;
            padding: 14px 36px;
            border-radius: 10px;
            font-size: 1rem;
            font-weight: 600;
            cursor: pointer;
            margin-top: 24px;
            transition: all 0.3s ease;
            display: inline-flex;
            align-items: center;
            gap: 8px;
        }

        .btn-download:hover {
            transform: translateY(-2px);
            box-shadow: 0 8px 25px rgba(16, 185, 129, 0.4);
        }

        .json-output {
            background: rgba(15, 23, 42, 0.9);
            border-radius: 12px;
            padding: 24px;
            margin-top: 28px;
            border: 1px solid var(--border);
        }

        .json-output label {
            text-transform: none;
            letter-spacing: 0;
            font-size: 0.9rem;
            margin-bottom: 16px;
            color: var(--text-secondary);
            display: block;
        }

        .json-output pre {
            color: var(--accent-light);
            font-family: 'Consolas', 'Monaco', monospace;
            font-size: 0.85rem;
            overflow-x: auto;
            white-space: pre-wrap;
            word-wrap: break-word;
            line-height: 1.6;
        }

        .bug-item {
            background: var(--bg-card);
            border-radius: 12px;
            padding: 24px;
            margin-bottom: 20px;
            border-left: 4px solid var(--accent);
            transition: transform 0.2s ease, box-shadow 0.2s ease;
        }

        .bug-item:hover {
            transform: translateX(4px);
            box-shadow: 0 4px 20px rgba(6, 182, 212, 0.1);
        }

        .bug-number {
            color: var(--accent-light);
            font-weight: 600;
            margin-bottom: 16px;
            font-size: 1.1rem;
            display: flex;
            align-items: center;
            gap: 8px;
        }

        .bug-number::before {
            content: '';
            width: 8px;
            height: 8px;
            background: var(--accent);
            border-radius: 50%;
        }

        @media (max-width: 768px) {
            body { padding: 20px 16px; }
            h1 { font-size: 2rem; }
            .card { padding: 24px; }
            .btn { padding: 16px 30px; }
        }
    </style>
</head>
<body>
    <div class="container">
        <div class="header">
            <h1>AI Bug Report Generator</h1>
            <p class="subtitle">Enter free-text bug descriptions (use numbering 1., 2., 3. for multiple bugs)</p>
        </div>

        <div class="card">
            <form id="bugForm">
                <div class="form-group">
                    <label for="description">Bug Description(s) *</label>
                    <textarea
                        name="description"
                        id="description"
                        placeholder="Enter one or multiple bugs...

Example:
1. When I try to login with wrong password, error message is not shown
2. Submit button freezes on the form page"
                    ></textarea>
                </div>

                <button type="submit" class="btn">Generate Bug Report</button>
            </form>

            <div class="warning-box" id="warningBox">
                <h3>⚠️ Need More Information</h3>
                <p id="warningMessage"></p>
            </div>

            <div class="result-box" id="resultBox">
                <h3>Generated Bug Reports</h3>
                <div id="reportContent"></div>
                <button class="btn-download" id="downloadBtn" onclick="downloadExcel()">
                    <svg width="18" height="18" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2" stroke-linecap="round" stroke-linejoin="round"><path d="M21 15v4a2 2 0 0 1-2 2H5a2 2 0 0 1-2-2v-4"/><polyline points="7 10 12 15 17 10"/><line x1="12" y1="15" x2="12" y2="3"/></svg>
                    Download Excel
                </button>
                <div class="json-output">
                    <label>JSON Output</label>
                    <pre id="jsonOutput"></pre>
                </div>
            </div>
        </div>
    </div>

    <script>
        let currentReport = null;

        document.getElementById('bugForm').addEventListener('submit', async (e) => {
            e.preventDefault();
            const desc = document.getElementById('description').value.trim();
            if (!desc) return alert('Please enter a bug description');

            const res = await fetch('/generate', {
                method: 'POST',
                headers: {'Content-Type': 'application/json'},
                body: JSON.stringify({ description: desc })
            });

            const data = await res.json();

            if (data.needs_clarification) {
                document.getElementById('warningBox').classList.add('show');
                document.getElementById('warningMessage').textContent = data.message;
                document.getElementById('resultBox').classList.remove('show');
                return;
            }

            currentReport = data;
            displayReport(data);
        });

        function displayReport(report) {
            document.getElementById('warningBox').classList.remove('show');
            let html = '';
            const fields = ['title', 'description', 'steps_to_reproduce', 'expected_behavior', 'actual_behavior', 'severity', 'priority', 'environment', 'test_type', 'status'];

            report.forEach((bug, index) => {
                html += '<div class="bug-item"><div class="bug-number">Bug #' + (index + 1) + '</div>';
                fields.forEach(f => {
                    const label = f.replace(/_/g, ' ');
                    html += '<div class="field"><div class="field-label">' + label + '</div><div class="field-value">' + (bug[f] || '') + '</div></div>';
                });
                html += '</div>';
            });

            document.getElementById('reportContent').innerHTML = html;
            document.getElementById('jsonOutput').textContent = JSON.stringify(report, null, 2);
            document.getElementById('resultBox').classList.add('show');
        }

        async function downloadExcel() {
            if (!currentReport) return;
            const res = await fetch('/download-excel', {
                method: 'POST',
                headers: {'Content-Type': 'application/json'},
                body: JSON.stringify(currentReport)
            });
            const blob = await res.blob();
            const url = window.URL.createObjectURL(blob);
            const a = document.createElement('a');
            a.href = url;
            a.download = 'bug_report_' + Date.now() + '.xlsx';
            a.click();
            window.URL.revokeObjectURL(url);
        }
    </script>
</body>
</html>'''

def split_bugs(description):
    text = description.strip()
    if not text:
        return []
    
    bugs = []
    
    lines = text.split('\n')
    
    for line in lines:
        line = line.strip()
        if not line:
            continue
        if re.match(r'^\d+[\.\)]\s+', line):
            if bugs and bugs[-1] not in ['SPLIT_MARKER']:
                bugs.append(line)
            else:
                bugs.append(line)
        elif bugs and bugs[-1] != 'SPLIT_MARKER':
            bugs[-1] = bugs[-1] + ' ' + line
        else:
            bugs.append(line)
    
    if not bugs:
        pattern = r'\d+[\.\)]\s+(?=[A-Z])'
        parts = re.split(pattern, text)
        bugs = [p.strip() for p in parts if p.strip()]
    
    if not bugs:
        bugs = [text]
    
    cleaned = []
    for bug in bugs:
        bug = bug.strip()
        bug = re.sub(r'^\d+[\.\)]\s+', '', bug)
        bug = bug.strip()
        if bug:
            cleaned.append(bug)
    
    return cleaned if cleaned else [text]

def generate_single_report(description):
    desc_lower = description.lower()
    
    severity_indicators = {'Critical': ['crash', 'freeze', 'data loss', 'security', 'breach'], 'Major': ['error', 'fail', 'not working', 'broken', 'wrong'], 'Medium': ['issue', 'problem', 'delay', 'slow'], 'Low': ['typo', 'cosmetic', 'minor']}
    severity = 'Medium'
    for sev, inds in severity_indicators.items():
        if any(i in desc_lower for i in inds):
            severity = sev
            break
    
    priority = {'Critical': 'High', 'Major': 'High', 'Medium': 'Medium', 'Low': 'Low'}[severity]
    
    environment = 'Web'
    if any(x in desc_lower for x in ['mobile', 'ios', 'android', 'app']): environment = 'Mobile'
    elif any(x in desc_lower for x in ['api', 'endpoint', 'rest', 'backend']): environment = 'API'
    
    test_type = 'Functional'
    if any(x in desc_lower for x in ['ui', 'visual', 'display', 'button', 'style']): test_type = 'UI'
    elif any(x in desc_lower for x in ['performance', 'slow', 'load', 'speed']): test_type = 'Performance'
    elif any(x in desc_lower for x in ['security', 'auth', 'login', 'password']): test_type = 'Security'
    
    action_words = ['click', 'enter', 'type', 'submit', 'select', 'navigate']
    found = [w for w in action_words if w in desc_lower]
    steps = "1. " + (found[0].title() if found else "Perform") + " the action\n2. Observe the behavior"
    
    words = description.split()
    title = ' '.join(words[:6])[:60] + ('...' if len(words) > 6 else '')
    
    expected = "System should work correctly"
    if 'error' in desc_lower or 'wrong' in desc_lower: expected = "Proper error message should be displayed"
    if 'login' in desc_lower: expected = "User should login or see appropriate error"
    
    return {
        "title": title,
        "description": description,
        "steps_to_reproduce": steps,
        "expected_behavior": expected,
        "actual_behavior": description,
        "severity": severity,
        "priority": priority,
        "environment": environment,
        "test_type": test_type,
        "status": "New"
    }

def generate_report(description):
    bug_descriptions = split_bugs(description)
    print(f"DEBUG: Found {len(bug_descriptions)} bugs: {bug_descriptions}")
    reports = []
    for bug_desc in bug_descriptions:
        if len(bug_desc.strip()) > 10:
            reports.append(generate_single_report(bug_desc))
    
    if not reports:
        return {"needs_clarification": True, "message": "Could not detect any valid bug descriptions. Please provide more details."}
    
    print(f"DEBUG: Generated {len(reports)} reports")
    return reports

def create_excel(reports):
    wb = Workbook()
    ws = wb.active
    ws.title = "Bug Reports"
    
    header_fill = PatternFill(start_color="00D9FF", end_color="00D9FF", fill_type="solid")
    header_font = Font(bold=True, color="000000")
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    
    severity_colors = {
        "critical": "FF0000",
        "major": "FF8C00",
        "high": "FF8C00",
        "medium": "FFD700",
        "minor": "FFFDD0",
        "low": "90EE90"
    }
    
    headers = ["#", "Title", "Description", "Steps to Reproduce", "Expected Behavior", "Actual Behavior", "Severity", "Priority", "Environment", "Test Type", "Status"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = thin_border
        cell.alignment = Alignment(horizontal='center')
    
    wrap_columns = ["Description", "Steps to Reproduce", "Expected Behavior", "Actual Behavior"]
    
    for row, report in enumerate(reports, 2):
        severity = report.get('severity', '').lower() if report.get('severity') else ''
        severity_fill = None
        for key, color in severity_colors.items():
            if key in severity:
                severity_fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
                break
        
        values = [row-1] + [report.get(f, '') for f in ['title', 'description', 'steps_to_reproduce', 'expected_behavior', 'actual_behavior', 'severity', 'priority', 'environment', 'test_type', 'status']]
        for col, value in enumerate(values, 1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.border = thin_border
            cell.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
            
            if col == 7 and severity_fill:
                cell.fill = severity_fill
        
        desc_val = report.get('description', '')
        steps_val = report.get('steps_to_reproduce', '')
        expected_val = report.get('expected_behavior', '')
        actual_val = report.get('actual_behavior', '')
        max_len = max(len(str(desc_val)), len(str(steps_val)), len(str(expected_val)), len(str(actual_val)))
        if max_len > 50:
            ws.row_dimensions[row].height = 80
        elif max_len > 30:
            ws.row_dimensions[row].height = 60
    
    ws.column_dimensions['A'].width = 6
    ws.column_dimensions['B'].width = 25
    ws.column_dimensions['C'].width = 35
    ws.column_dimensions['D'].width = 30
    ws.column_dimensions['E'].width = 30
    ws.column_dimensions['F'].width = 30
    ws.column_dimensions['G'].width = 12
    ws.column_dimensions['H'].width = 10
    ws.column_dimensions['I'].width = 12
    ws.column_dimensions['J'].width = 12
    ws.column_dimensions['K'].width = 10
    
    return wb

class Handler(BaseHTTPRequestHandler):
    def log_message(self, format, *args):
        print(f"{self.address_string()} - {format % args}")
    
    def do_GET(self):
        self.send_response(200)
        self.send_header('Content-type', 'text/html')
        self.end_headers()
        self.wfile.write(HTML.encode())
    
    def do_POST(self):
        try:
            if self.path == '/generate':
                length = int(self.headers['Content-Length'])
                data = self.rfile.read(length)
                desc = json.loads(data)['description']
                report = generate_report(desc)
                self.send_response(200)
                self.send_header('Content-type', 'application/json')
                self.end_headers()
                self.wfile.write(json.dumps(report).encode())
            elif self.path == '/download-excel':
                length = int(self.headers['Content-Length'])
                data = self.rfile.read(length)
                reports = json.loads(data)
                wb = create_excel(reports)
                import io
                buffer = io.BytesIO()
                wb.save(buffer)
                buffer.seek(0)
                self.send_response(200)
                self.send_header('Content-type', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
                self.send_header('Content-Disposition', 'attachment; filename=bug_report.xlsx')
                self.end_headers()
                self.wfile.write(buffer.read())
            else:
                self.send_response(404)
                self.end_headers()
        except Exception as e:
            print(f"ERROR: {e}")
            import traceback
            traceback.print_exc()
            self.send_response(500)
            self.send_header('Content-type', 'application/json')
            self.end_headers()
            self.wfile.write(json.dumps({"error": str(e)}).encode())

print("Starting server at http://127.0.0.1:5000")
server = HTTPServer(('127.0.0.1', 5000), Handler)
server.serve_forever()